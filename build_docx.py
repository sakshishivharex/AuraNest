# build_docx.py — merge the report markdown files into one formatted .docx.
#
#   python build_docx.py
#
# Output: docs/AuraNest MCA Project Report - Sakshi Shivhare.docx
# Formatting follows README.md: Times New Roman 12pt body, 16/14/12pt bold
# headings, justified, 1.15 spacing, normal margins, page numbers bottom centre.

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

DOCS = Path(__file__).parent
ORDER = [
    "01-front-matter.md",
    "02-chapter-1-introduction.md",
    "03-chapter-2-literature-review.md",
    "04-chapter-3-system-analysis-design.md",
    "05-chapter-4-implementation-results.md",
    "06-chapter-5-conclusion.md",
    "07-references.md",
    "08-appendix.md",
]

# Reports where every level-2 heading starts a new page (title page = page 1).
BREAK_H2_FILES = {"frontend/report.md", "backend/report.md",
                  "database/report.md"}

INLINE = re.compile(r"(\*\*.+?\*\*|(?<!\*)\*[^*]+?\*(?!\*)|`[^`]+?`)")


# ----------------------------------------------------------------- helpers

def set_font(run, name="Times New Roman", size=12, bold=None, italic=None,
             color=None, mono=False):
    if mono:
        run.font.name = "Courier New"
        rpr = run._element.get_or_add_rPr()
        rfonts = rpr.find(qn("w:rFonts"))
        if rfonts is None:
            rfonts = OxmlElement("w:rFonts")
            rpr.append(rfonts)
        for attr in ("w:ascii", "w:hAnsi", "w:cs"):
            rfonts.set(qn(attr), "Courier New")
    else:
        run.font.name = name
        rpr = run._element.get_or_add_rPr()
        rfonts = rpr.find(qn("w:rFonts"))
        if rfonts is None:
            rfonts = OxmlElement("w:rFonts")
            rpr.append(rfonts)
        for attr in ("w:ascii", "w:hAnsi", "w:cs"):
            rfonts.set(qn(attr), name)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = color


def add_inline(paragraph, text, size=12, base_bold=False, base_italic=False,
               mono=False):
    """Parse **bold**, *italic* and `code` spans into runs."""
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            run = paragraph.add_run(part[2:-2])
            set_font(run, size=size, bold=True, italic=base_italic, mono=mono)
        elif part.startswith("`") and part.endswith("`") and len(part) > 2:
            run = paragraph.add_run(part[1:-1])
            set_font(run, name="Courier New" if not mono else "Courier New",
                     size=size - 1 if not mono else size, mono=True)
            shd = OxmlElement("w:shd")
            shd.set(qn("w:val"), "clear")
            shd.set(qn("w:fill"), "F2F2F2")
            run._element.get_or_add_rPr().append(shd)
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            run = paragraph.add_run(part[1:-1])
            set_font(run, size=size, bold=base_bold, italic=True, mono=mono)
        else:
            run = paragraph.add_run(part)
            set_font(run, size=size, bold=base_bold, italic=base_italic,
                     mono=mono)


def shade(element, fill="F2F2F2"):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), fill)
    element.append(shd)


def page_break(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    p.add_run().add_break(WD_BREAK.PAGE)


def add_page_number(section):
    footer = section.footer
    p = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    set_font(run, size=11)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(end)


def heading(doc, text, level, page_break_before):
    if page_break_before:
        page_break(doc)
    sizes = {1: 16, 2: 14, 3: 12, 4: 12}
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    pf = p.paragraph_format
    pf.space_before = Pt(0 if page_break_before else 14)
    pf.space_after = Pt(8)
    pf.keep_with_next = True
    pf.line_spacing = 1.15
    add_inline(p, text, size=sizes[level], base_bold=True,
               base_italic=(level >= 4))
    return p


def is_table_sep(line):
    return line.startswith("|") and set(line.replace("|", "").replace(
        ":", "").replace("-", "").replace(" ", "")) == set()


def split_row(line):
    return [c.strip() for c in line.strip().strip("|").split("|")]


def add_table(doc, lines):
    rows = [split_row(l) for l in lines if not is_table_sep(l)]
    if not rows:
        return
    cols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=cols)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row in enumerate(rows):
        for j in range(cols):
            cell = table.cell(i, j)
            text = row[j] if j < len(row) else ""
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.0
            p.alignment = (WD_ALIGN_PARAGRAPH.CENTER if i == 0
                           else WD_ALIGN_PARAGRAPH.LEFT)
            add_inline(p, text, size=10.5, base_bold=(i == 0))
            if i == 0:
                shade(cell._tc.get_or_add_tcPr(), "E7E6E6")
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def add_code(doc, code):
    for i, line in enumerate(code.splitlines()):
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.space_before = Pt(4 if i == 0 else 0)
        pf.space_after = Pt(4 if i == len(code.splitlines()) - 1 else 0)
        pf.line_spacing = 1.0
        pf.left_indent = Cm(0.4)
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run = p.add_run(line if line else " ")
        set_font(run, size=9, mono=True)
        shade(p._p.get_or_add_pPr(), "F5F5F5")


# ------------------------------------------------------------------ parser

def convert(doc, text, filename):
    lines = text.splitlines()
    if filename == "01-front-matter.md":
        # drop the split-file marker heading; the title page starts page 1
        lines = [l for l in lines if l.strip() != "# Front Matter"]
    lines = [l for l in lines if l.strip() not in ("", "---", "<center>",
                                                   "</center>")]

    i = 0
    first_h2_done = False
    first_h1_seen = False
    while i < len(lines):
        line = lines[i]

        # fenced code
        if line.startswith("```"):
            j = i + 1
            buf = []
            while j < len(lines) and not lines[j].startswith("```"):
                buf.append(lines[j])
                j += 1
            add_code(doc, "\n".join(buf))
            i = j + 1
            continue

        # headings
        m = re.match(r"^(#{1,4})\s+(.*)$", line)
        if m:
            level = len(m.group(1))
            title = m.group(2).strip()
            if level == 1:
                pb = first_h1_seen
                first_h1_seen = True
            elif level == 2 and (filename == "01-front-matter.md"
                                 or filename in BREAK_H2_FILES):
                pb = first_h2_done
                first_h2_done = True
            else:
                pb = False
            heading(doc, title, level, pb)
            i += 1
            continue

        # table
        if line.startswith("|") and i + 1 < len(lines) and is_table_sep(
                lines[i + 1]):
            j = i
            buf = []
            while j < len(lines) and lines[j].startswith("|"):
                buf.append(lines[j])
                j += 1
            add_table(doc, buf)
            i = j
            continue

        # blockquote
        if line.startswith(">"):
            p = doc.add_paragraph()
            pf = p.paragraph_format
            pf.left_indent = Cm(0.8)
            pf.line_spacing = 1.15
            pf.space_after = Pt(6)
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            add_inline(p, line.lstrip("> ").strip(), size=12,
                       base_italic=True)
            i += 1
            continue

        # numbered list
        m = re.match(r"^(\d+)\.\s+(.*)$", line)
        if m:
            p = doc.add_paragraph(style="List Number")
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.space_after = Pt(4)
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            add_inline(p, m.group(2), size=12)
            i += 1
            continue

        # bullet list
        if line.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.space_after = Pt(4)
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            add_inline(p, line[2:].strip(), size=12)
            i += 1
            continue

        # plain paragraph
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.space_after = Pt(8)
        add_inline(p, line.strip(), size=12)
        i += 1


# ------------------------------------------------------------------- main

def main():
    import sys

    # Modes:
    #   python build_docx.py                      -> combined report (01..08)
    #   python build_docx.py frontend/report.md   -> one split report
    args = sys.argv[1:]

    doc = Document()

    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    rpr = normal.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs"):
        rfonts.set(qn(attr), "Times New Roman")

    for section in doc.sections:
        section.top_margin = Cm(2.54)
        section.bottom_margin = Cm(2.54)
        section.left_margin = Cm(2.54)
        section.right_margin = Cm(2.54)
        add_page_number(section)

    if args:
        order = args
    else:
        order = ORDER
    for filename in order:
        path = DOCS / filename
        if not path.exists():
            raise SystemExit(f"missing: {path}")
        convert(doc, path.read_text(encoding="utf-8"), filename)

    props = doc.core_properties
    props.title = ("AuraNest - A Full-Stack E-Commerce Web Application "
                   "for Home Furnishings")
    props.author = "Sakshi Shivhare"
    props.subject = "MCA Project Report, Chandigarh University"

    if args and len(args) == 1:
        label = Path(args[0]).parent.name
        out = DOCS / f"AuraNest {label.capitalize()} Report - Sakshi Shivhare.docx"
    else:
        out = DOCS / "AuraNest MCA Project Report - Sakshi Shivhare.docx"
    doc.save(out)
    print(f"written: {out}")


if __name__ == "__main__":
    main()
